"""Recalculate a workbook with headless LibreOffice and scan it for formula errors.

    python tests/lo_tools.py Operations_Dashboard.xlsx            # recalc a copy, report errors
    python tests/lo_tools.py Operations_Dashboard.xlsx --in-place

openpyxl writes formulas without cached values, so LibreOffice computes every
formula and saves the results; the scan then reads those cached values.
"""
from __future__ import annotations

import argparse
import json
import os
import shutil
import subprocess
import sys
import tempfile
import time
from collections import Counter
from pathlib import Path

from openpyxl import load_workbook

ERRORS = ("#N/A", "#DIV/0!", "#VALUE!", "#REF!", "#NAME?", "#NUM!", "#NULL!", "Err:")

MACRO = """<?xml version="1.0" encoding="UTF-8"?>
<!DOCTYPE script:module PUBLIC "-//OpenOffice.org//DTD OfficeDocument 1.0//EN" "module.dtd">
<script:module xmlns:script="http://openoffice.org/2000/script" script:name="Module1" script:language="StarBasic">
Sub RecalculateAndSave(path As String)
  Dim args(0) As New com.sun.star.beans.PropertyValue
  args(0).Name = "Hidden"
  args(0).Value = True
  doc = StarDesktop.loadComponentFromURL(ConvertToURL(path), "_blank", 0, args())
  doc.calculateAll()
  doc.store()
  doc.close(True)
End Sub
</script:module>"""


def recalc(path: Path, timeout: int = 3600) -> float:
    """Recalculate `path` in place; returns seconds taken."""
    profile = Path(tempfile.mkdtemp(prefix="lo_profile_"))
    env = dict(os.environ, SAL_USE_VCLPLUGIN="svp")
    url = profile.as_uri()
    subprocess.run(["soffice", "--headless", "--terminate_after_init", f"-env:UserInstallation={url}"],
                   env=env, capture_output=True, timeout=180)
    basic = profile / "user" / "basic" / "Standard"
    basic.mkdir(parents=True, exist_ok=True)
    (basic / "Module1.xba").write_text(MACRO)
    t0 = time.time()
    macro = f'macro:///Standard.Module1.RecalculateAndSave("{path.resolve()}")'
    subprocess.run(["soffice", "--headless", "--norestore", f"-env:UserInstallation={url}", macro],
                   env=env, capture_output=True, timeout=timeout)
    shutil.rmtree(profile, ignore_errors=True)
    return time.time() - t0


def scan(path: Path) -> dict:
    """Count error values in every sheet (reads cached values)."""
    wb = load_workbook(path, read_only=True, data_only=True)
    found, where = Counter(), {}
    for ws in wb.worksheets:
        for row in ws.iter_rows():
            for c in row:
                v = c.value
                if isinstance(v, str) and v.startswith(ERRORS):
                    key = v.split()[0]
                    found[key] += 1
                    where.setdefault(key, []).append(f"{ws.title}!{c.coordinate}")
    wb.close()
    return {"total_errors": sum(found.values()), "by_type": dict(found),
            "examples": {k: v[:20] for k, v in where.items()}}


def values(path: Path, sheet: str, cells: list[str]) -> dict:
    wb = load_workbook(path, data_only=True, read_only=False)
    ws = wb[sheet]
    out = {c: ws[c].value for c in cells}
    wb.close()
    return out


def main(argv=None):
    ap = argparse.ArgumentParser()
    ap.add_argument("workbook")
    ap.add_argument("--in-place", action="store_true")
    args = ap.parse_args(argv)
    src = Path(args.workbook)
    target = src if args.in_place else src.with_name(src.stem + "_recalc.xlsx")
    if target != src:
        shutil.copy(src, target)
    secs = recalc(target)
    result = scan(target)
    result["recalc_seconds"] = round(secs, 1)
    result["file"] = str(target)
    print(json.dumps(result, indent=2))
    return 1 if result["total_errors"] else 0


if __name__ == "__main__":
    sys.exit(main())
