"""End-to-end tests: generate data -> ingest -> build -> LibreOffice recalc -> check numbers.

Run with:  python -m pytest tests -q        (needs LibreOffice for the workbook tests)

The expected values are recomputed independently with pandas from the
databank and compared with what the workbook formulas produce.
"""
from __future__ import annotations

import datetime as dt
import shutil
import sqlite3
import sys
from pathlib import Path

import pandas as pd
import pytest
from openpyxl import load_workbook

HERE = Path(__file__).resolve().parent
sys.path.insert(0, str(HERE.parent))
sys.path.insert(0, str(HERE))

import build_workbook  # noqa: E402
import generate_sample_data  # noqa: E402
import ingest  # noqa: E402
from lo_tools import recalc, scan  # noqa: E402

TODAY = dt.date(2026, 9, 28)
N_WEEKS = 16
HAS_LO = shutil.which("soffice") is not None


@pytest.fixture(scope="module")
def root(tmp_path_factory):
    r = tmp_path_factory.mktemp("ops")
    generate_sample_data.generate(r, n_weeks=N_WEEKS, seed=11, today=TODAY, scale=0.1)
    ingest.run(r, build=False, quiet=True)
    build_workbook.build(r)
    return r


def db(root) -> sqlite3.Connection:
    return sqlite3.connect(root / "databank.sqlite")


# --------------------------------------------------------------------------
# Ingestion
# --------------------------------------------------------------------------


def test_all_streams_loaded(root):
    con = db(root)
    for table in ("wsf_cells", "wsf_zone", "midweek", "chop", "community", "transport_runs", "transport_costs",
                  "transport_budget"):
        weeks = con.execute(f"SELECT COUNT(DISTINCT week_ending) FROM {table}").fetchone()[0]
        assert weeks == N_WEEKS, table


def test_injected_errors_are_logged(root):
    issues = pd.read_sql("SELECT stream, severity, issue FROM dq_log", db(root))
    got = set(zip(issues.stream, issues.issue))
    for expected in [("WSF_CELLS", "Grand total mismatch"), ("WSF_CELLS", "male is negative"),
                     ("WSF_CELLS", "Unknown Cell (not in MASTER_CELLS)"), ("WSF_CELLS", "Duplicate row in file"),
                     ("WSF_CELLS", "Date missing or not recognised"),
                     ("MIDWEEK", "Community Church row in a zonal file (kept separate by design)"),
                     ("MIDWEEK", "Unknown Zone (not in MASTER_ZONES)"),
                     ("CHOP", "CHOP report from a zone without CHOP (Has_CHOP = N)"),
                     ("CHOP", "Adult total mismatch"),
                     ("TRANSPORT FT_PROCURED", "Unknown Area"),
                     ("TRANSPORT FT_PROCURED", "Number cleaned"),
                     ("TRANSPORT COASTER_FUEL", "Pump price differs from the rest of the sheet")]:
        assert expected in got, expected


def test_chop_only_from_chop_zones(root):
    con = db(root)
    zones = pd.read_csv(root / "masters" / "MASTER_ZONES.csv", dtype=str)
    chop_zones = set(zones.loc[zones.Has_CHOP == "Y", "Zone_Name"])
    loaded = {z for (z,) in con.execute("SELECT DISTINCT zone FROM chop")}
    assert loaded <= chop_zones


def test_community_never_in_zonal_tables(root):
    con = db(root)
    churches = set(pd.read_csv(root / "masters" / "MASTER_COMMUNITY.csv", dtype=str)["Church_Name"])
    for table in ("wsf_zone", "midweek", "chop"):
        zones = {z for (z,) in con.execute(f"SELECT DISTINCT zone FROM {table}")}
        assert not zones & churches, table


def test_wsf_zone_rows_cover_every_active_zone(root):
    con = db(root)
    n_active = len(pd.read_csv(root / "masters" / "MASTER_ZONES.csv"))
    per_week = con.execute("SELECT COUNT(*) FROM wsf_zone GROUP BY week_ending").fetchall()
    assert {n for (n,) in per_week} == {n_active}
    agg = con.execute("SELECT SUM(male), SUM(female), SUM(children) FROM wsf_zone").fetchone()
    raw = con.execute("SELECT SUM(male), SUM(female), SUM(children) FROM wsf_cells").fetchone()
    assert agg == raw


def test_reingest_same_file_does_not_double_count(root, tmp_path):
    con = db(root)
    before = {t: con.execute(f"SELECT COUNT(*), SUM(male) FROM {t}").fetchone()
              for t in ("wsf_cells", "wsf_zone", "midweek", "chop")}
    transport_before = {t: con.execute(f"SELECT COUNT(*) FROM {t}").fetchone()
                        for t in ("transport_runs", "transport_costs", "transport_budget")}
    con.close()
    processed = root / "inbox" / "processed"
    resend = [p for p in processed.iterdir() if "Midweek" in p.name or "WSF cell" in p.name
              or "Covenant" in p.name][:6]
    resend += [p for p in processed.iterdir() if "Transport raw" in p.name or "FUELING" in p.name
               or "FT PROCURED" in p.name][:3]
    for p in resend:
        shutil.copy(p, root / "inbox" / p.name.split("_", 1)[1])
    summary = ingest.run(root, build=False, quiet=True)
    con = db(root)
    after = {t: con.execute(f"SELECT COUNT(*), SUM(male) FROM {t}").fetchone() for t in before}
    assert after == before
    for t, before_t in transport_before.items():
        assert con.execute(f"SELECT COUNT(*) FROM {t}").fetchone() == before_t, t
    # everything re-sent was counted as a replacement, not a new row
    for stream, (loaded, replaced, _rej) in summary.items():
        assert loaded == replaced, stream


def test_header_detection_and_aliases(tmp_path):
    aliases = ingest.load_aliases()
    raw = pd.DataFrame([["MIDWEEK RETURNS", None, None, None, None], [None] * 5,
                        ["No. of Females", "zone name", "KIDS", "Meeting-Date", "MEN"],
                        [10, "X Zone 01", 3, "02/09/2026", 7]])
    hdr, mapping = ingest.find_header(raw, aliases)
    assert hdr == 2
    assert sorted(mapping.values()) == ["children", "date", "female", "male", "zone"]
    assert ingest.detect_stream(set(mapping.values()), "Zonal Midweek wk36.xlsx", "Sheet1") == "MIDWEEK"
    assert ingest.parse_date("02/09/2026") == dt.date(2026, 9, 2)   # day first (Nigeria)


SAMPLES = Path(__import__("os").environ.get(
    "OPS_SAMPLES", "/root/.claude/uploads/e0198d9f-6c86-5d5b-85b9-0a9576ba4c38"))


@pytest.mark.skipif(not list(SAMPLES.glob("*Transport_raw_reports*.xlsx")), reason="real sample files not present")
def test_real_transport_samples_27_sept(tmp_path):
    """The transport office's own files for 27 Sept 2026 (not in the repo: they contain names)."""
    root = tmp_path / "real"
    (root / "masters").mkdir(parents=True)
    shutil.copy(HERE.parent / "masters" / "MASTER_AREAS.csv", root / "masters" / "MASTER_AREAS.csv")
    (root / "inbox").mkdir()
    for pattern in ("*Transport_raw_reports*", "*FT_COASTERS_FUELING*", "*FT_PROCURED_FOR*", "*EV__Tata*",
                    "*ZONE_DETAILS*"):
        for f in list(SAMPLES.glob(pattern))[:1]:
            shutil.copy(f, root / "inbox" / f.name)
    ingest.run(root, build=False, quiet=True)
    con = db(root)
    q = lambda sql: con.execute(sql).fetchone()
    assert q("SELECT DISTINCT week_ending FROM transport_runs") == ("2026-09-27",)
    # WSF Procured matches the sheet's own riders and cost totals
    assert q("SELECT SUM(male+female+children), SUM(cost) FROM transport_runs WHERE category='WSF Procured'") == (
        9775, 19670000.0)
    # FT Procured: every bus row is a bus (a blank "in church" cell is only a missing loading-bay entry);
    # cost is the sheet's ₦54,267,500 plus Ikoyi's text '170-,000'
    assert q("SELECT COUNT(*), SUM(cost) FROM transport_runs WHERE category='FT Procured'") == (1019, 54437500.0)
    # coasters marked ZONE on the fuel schedule load at loading bays; TATA/EV at hubs
    assert q("SELECT COUNT(*) FROM transport_runs WHERE category='Church Coaster' AND location_type='Loading Bay'") == (46,)
    assert q("SELECT COUNT(*) FROM transport_runs WHERE report='EV_TATA_REPORT' AND location_type<>'Hub'") == (0,)
    assert q("SELECT severity FROM dq_log WHERE issue LIKE 'Hub Payment Summary%'") == ("Info",)
    if list(SAMPLES.glob("*ZONE_DETAILS*")):
        # the zone list fills loading-bay addresses the FT sheet left blank, and WSF Procured locations
        assert q("SELECT COUNT(*) FROM transport_runs WHERE category='FT Procured' AND location<>''")[0] >= 955
        assert q("SELECT COUNT(*) FROM transport_runs WHERE category='WSF Procured' AND location<>''")[0] >= 274
    # coaster fuel actually paid (HUB coasters) = ₦7,574,800
    assert q("SELECT SUM(amount) FROM transport_costs WHERE payable='Y'") == (7574800.0,)
    assert q("SELECT COUNT(*) FROM transport_budget") == (92,)
    assert q("SELECT COUNT(*) FROM dq_log WHERE severity='Rejected'") == (0,)


# --------------------------------------------------------------------------
# Workbook scenarios (need LibreOffice)
# --------------------------------------------------------------------------

def expected_wsf(root, scope: str, sel: str, week: dt.date | None, n: int) -> dict:
    z = pd.read_sql("SELECT * FROM wsf_zone", db(root))
    z["week"] = pd.to_datetime(z.week_ending).dt.date
    if scope == "Area":
        z = z[z.area == sel]
    elif scope == "Zone":
        z = z[z.zone == sel]
    z["grand"] = z.male + z.female + z.children
    wk = z.groupby("week").agg(grand=("grand", "sum"), cells=("cells_total", "sum"), rep=("cells_reported", "sum"))
    latest = max(pd.read_sql("SELECT week_ending FROM wsf_zone", db(root)).week_ending.map(dt.date.fromisoformat))
    sel_w = week or latest
    prev = sel_w - dt.timedelta(days=7)
    window = [sel_w - dt.timedelta(days=7 * k) for k in range(1, n + 1)]
    win = wk[wk.index.isin(window)]
    mtd = wk[[(w.month == sel_w.month and w.year == sel_w.year and w <= sel_w) for w in wk.index]]
    ytd = wk[[(w.year == sel_w.year and w <= sel_w) for w in wk.index]]
    return {"grand": wk.grand.get(sel_w), "prev": wk.grand.get(prev),
            "navg": win.grand.mean() if len(win) else None, "mtd": mtd.grand.sum(), "ytd": ytd.grand.sum(),
            "pct": wk.rep.get(sel_w) / wk.cells.get(sel_w),
            "pct_navg": win.rep.sum() / win.cells.sum() if len(win) else None,
            "week": sel_w}


def expected_nonreporting(root, scope, sel, week) -> int:
    con = db(root)
    cells = pd.read_csv(root / "masters" / "MASTER_CELLS.csv", dtype=str)
    zones = pd.read_csv(root / "masters" / "MASTER_ZONES.csv", dtype=str)
    cells = cells[cells.Operational == "Y"].merge(zones[["Zone_Name", "Area_Name"]], on="Zone_Name")
    rep = {c for (c,) in con.execute("SELECT cell FROM wsf_cells WHERE week_ending = ?", (week.isoformat(),))}
    if scope == "Area":
        cells = cells[cells.Area_Name == sel]
    elif scope == "Zone":
        cells = cells[cells.Zone_Name == sel]
    return int((~cells.Cell.isin(rep)).sum())


def run_scenario(root, tmp_path, name, scope="Global", select=None, week=None, window=4):
    wb = load_workbook(root / "Operations_Dashboard.xlsx")
    ws = wb["WSF DASHBOARD"]
    ws["B4"] = scope
    ws["E4"] = select
    ws["I4"] = week if week else "Latest"
    ws["K4"] = window
    out = tmp_path / f"{name}.xlsx"
    wb.save(out)
    recalc(out)
    return out


def calc_values(path):
    wb = load_workbook(path, data_only=True)
    c, d = wb["CALC_WSF"], wb["WSF DASHBOARD"]
    vals = {"grand": c["B90"].value, "prev": c["C90"].value, "navg": c["F90"].value, "mtd": c["H90"].value,
            "ytd": c["J90"].value, "pct": c["B98"].value, "pct_navg": c["F98"].value, "week": c["B10"].value,
            "nonrep": c["AX1"].value, "rank_head": d["B46"].value, "card": d["B8"].value, "status": d["B5"].value,
            "rank_box": d["O46"].value, "top_area": c["AJ3"].value}
    wb.close()
    return vals


def close(a, b, tol=1e-6):
    if b is None or (isinstance(b, float) and pd.isna(b)):
        return a in ("", None)
    return a not in ("", None) and abs(float(a) - float(b)) <= tol * max(1, abs(float(b)))


SCENARIOS = [
    ("global_w4", "Global", None, None, 4),
    ("global_w1", "Global", None, None, 1),
    ("global_w13", "Global", None, None, 13),
    ("area_w4", "Area", "AREA", None, 4),
    ("zone_chop_w13", "Zone", "ZONE_CHOP", None, 13),
    ("zone_nochop_w4", "Zone", "ZONE_NOCHOP", None, 4),
    ("older_week_w4", "Global", None, "PREV3", 4),
]


@pytest.mark.skipif(not HAS_LO, reason="LibreOffice not installed")
@pytest.mark.parametrize("name,scope,select,week,window", SCENARIOS)
def test_wsf_dashboard_scenarios(root, tmp_path, name, scope, select, week, window):
    zones = pd.read_csv(root / "masters" / "MASTER_ZONES.csv", dtype=str)
    pick = {"AREA": zones.Area_Name.iloc[len(zones) // 2],
            "ZONE_CHOP": zones[zones.Has_CHOP == "Y"].Zone_Name.iloc[3],
            "ZONE_NOCHOP": zones[zones.Has_CHOP == "N"].Zone_Name.iloc[3]}
    sel = pick.get(select, select)
    latest = max(pd.read_sql("SELECT week_ending FROM wsf_zone", db(root)).week_ending.map(dt.date.fromisoformat))
    wk = latest - dt.timedelta(weeks=3) if week == "PREV3" else None
    out = run_scenario(root, tmp_path, name, scope, sel, wk, window)
    got = calc_values(out)
    exp = expected_wsf(root, scope, sel, wk, window)
    for k in ("grand", "prev", "navg", "mtd", "ytd", "pct", "pct_navg"):
        assert close(got[k], exp[k]), (k, got[k], exp[k])
    assert got["week"].date() == exp["week"]
    if (latest - exp["week"]).days < 7 * 8:
        assert got["nonrep"] == expected_nonreporting(root, scope, sel, exp["week"])
    if scope == "Zone":
        assert got["rank_head"] == "n/a at Zone scope"
    else:
        assert got["rank_head"].startswith("1. ")
    if scope == "Area":
        assert got["rank_box"] != "n/a"
    assert scan(out)["total_errors"] == 0


@pytest.mark.skipif(not HAS_LO, reason="LibreOffice not installed")
@pytest.mark.parametrize("scope", ["Global", "Area"])
def test_transport_dashboard(root, tmp_path, scope):
    """Transport only: dashboard numbers match an independent pandas calculation."""
    areas = pd.read_csv(root / "masters" / "MASTER_AREAS.csv", dtype=str)
    sel = areas.Area_Name.iloc[1] if scope == "Area" else None
    wb = load_workbook(root / "Operations_Dashboard.xlsx")
    ws = wb["TRANSPORT DASHBOARD"]
    ws["B4"], ws["E4"], ws["I4"], ws["K4"] = scope, sel, "Latest", 4
    out = tmp_path / f"transport_{scope}.xlsx"
    wb.save(out)
    recalc(out)
    assert scan(out)["total_errors"] == 0
    c = load_workbook(out, data_only=True)["CALC_TRANSPORT"]
    con = db(root)
    runs = pd.read_sql("SELECT * FROM transport_runs", con)
    costs = pd.read_sql("SELECT * FROM transport_costs", con)
    budget = pd.read_sql("SELECT * FROM transport_budget", con)
    wk = runs.week_ending.max()
    if sel:
        runs, costs, budget = runs[runs.area == sel], costs[costs.area == sel], budget[budget.area == sel]
    r, cs, b = runs[runs.week_ending == wk], costs[costs.week_ending == wk], budget[budget.week_ending == wk]
    riders = int((r.male + r.female + r.children).sum())
    spend = r.cost.sum() + cs.loc[cs.payable == "Y", "amount"].sum()
    ran = r.loc[r.status == "Ran", "buses"].sum()
    seats_rows = r[(r.status == "Ran") & r.capacity.notna() & (r.capacity > 0)]
    seats = (seats_rows.capacity * seats_rows.trips.clip(lower=1)).sum()
    seat_cap = seats_rows.capacity * seats_rows.trips.clip(lower=1)
    riders_on = seats_rows.male + seats_rows.female + seats_rows.children
    util = riders_on.clip(upper=seat_cap).sum() / seats          # a full bus counts as 100%
    excess = (riders_on - seat_cap).clip(lower=0).sum()
    ft_budget = (b.buses_allocated * b.cost_per_bus).sum()
    ft_cost = r.loc[r.category == "FT Procured", "cost"].sum()
    got = {k: c[f"B{90 + i}"].value for i, k in enumerate(
        ["riders", "ran", "listed", "pct_ran", "seats", "util", "excess", "spend", "central", "members", "fuel", "cpr", "cpb",
         "cps", "ft_vs_budget", "breakdowns"])}
    assert got["riders"] == riders
    assert got["ran"] == ran
    assert abs(got["spend"] - spend) < 0.01
    assert abs(got["seats"] - seats) < 0.01
    assert abs(got["util"] - util) < 1e-9
    assert abs(got["excess"] - excess) < 1e-9
    assert abs(got["ft_vs_budget"] - ft_cost / ft_budget) < 1e-9
    assert got["breakdowns"] == int((r.status == "Breakdown").sum())


@pytest.mark.skipif(not HAS_LO, reason="LibreOffice not installed")
def test_empty_workbook_has_no_errors(root, tmp_path):
    out = tmp_path / "empty.xlsx"
    build_workbook.build(root, out, empty=True)
    recalc(out)
    assert scan(out)["total_errors"] == 0
    got = calc_values(out)
    assert got["card"] == "–"
    assert got["status"].startswith("No WSF data loaded yet")


def test_zone_master_rules(tmp_path):
    """ZONE DETAILS: codes normalised and corrected to their Area, duplicates and clashes handled."""
    import zone_master
    from databank import Masters
    (tmp_path / "masters").mkdir()
    pd.DataFrame([{"Area_ID": "A01", "Area_Name": "CANAANLAND 2", "Active": "Y", "Area_No": "1", "Aliases": ""},
                  {"Area_ID": "A87", "Area_Name": "ILOGBO OKOKO", "Active": "Y", "Area_No": "87",
                   "Aliases": "ILOGBO-OKOKO"}]).to_csv(tmp_path / "masters" / "MASTER_AREAS.csv", index=False)
    m = Masters.load(tmp_path / "masters")
    zones = pd.DataFrame([["area number", "area name", "Zone", "Zin Codes"],
                          [1, "CANAANLAND 2", "Zone 1 - 51/53 IGE DARAMOLA", "LFC 0101"],
                          [1, "CANAANLAND 2", "Zone 2 - 9 BENJA ROAD, OTA", "lfcO102"],
                          [1, "CANAANLAND 2", "Zone 2 - 9 BENJA ROAD, OTA", "LFC0102"],      # listed twice
                          [1, "CANAANLAND 2", "Zone 3 - 1 JESU O SEUN STREET", None],         # no code
                          [87, "ILOGBO-OKOKO", "Zone 14 - JESAB ZONE", "LFC6714"]])          # wrong Area digits
    new = pd.DataFrame([["AREA NAME(YABA)", "AREA NO(15)", "ZONAL ADDRESS", "PROPOSED ZONAL NO (LFC1511)"],
                        ["CANAANLAND 2", 1, "9 BENJA ROAD OTA", "LFC0102"],                  # already listed
                        ["CANAANLAND 2", 1, "5 NEW ROAD, OTA", "LFC0101"],                   # clashes
                        ["CANAANLAND 2", 1, "7 FRESH STREET", "LFC0104"]])
    df, issues = zone_master.build({"ZONES": zones, "NEW ZONES": new}, m, m.zones, {"LFC0199": "CANAANLAND 2"})
    got = dict(zip(df.Zone_ID, df.Zone_Status))
    assert got == {"LFC0101": "Existing", "LFC0102": "Existing", "LFC0103": "Existing", "LFC0104": "New (proposed)",
                   "LFC0199": "Seen in reports only", "LFC8714": "Existing"}
    assert df.Zone_ID.is_unique
    kinds = {i["issue"] for i in issues}
    assert {"Zone code corrected to its Area", "Zone listed twice", "Zone code inferred", "New zone already listed",
            "Proposed code already in use", "Zone code not in ZONE DETAILS"} <= kinds
    df.to_csv(tmp_path / "masters" / "MASTER_ZONES.csv", index=False)
    m2 = Masters.load(tmp_path / "masters")
    assert m2.zone_for_address("9, Benja Road, Ota", "CANAANLAND 2") == "LFC0102"
    assert m2.zone_address["LFC8714"] == "JESAB ZONE"
