"""The control strip shared by every dashboard, and the helper ranges behind it.

Controls (row 4 of every dashboard):
    B4  Scope        Global / Area / Zone
    C4  Search       free text; part of an Area or Zone name
    E4  Select       drop-down that lists only names matching Search
    H4  Stream       fixed label on stream dashboards, a selector on the Master
    I4  Week ending  "Latest" or a week that has data
    K4  Window       1, 2, 4, 8, 13, 26 or 52 weeks

The search drop-down is built without dynamic arrays (FILTER/SORT) so it
works in every Excel version and in LibreOffice: a helper column marks the
rows whose name contains the search text, SMALL() lists them in order, and a
defined name sized with the match count feeds the data-validation list.
"""
from __future__ import annotations

from openpyxl.worksheet.datavalidation import DataValidation

from .common import BOX, CENTER, LEFT, abs_ref, col, define, fill, font, put, tint

SCOPE, SEARCH, SELECT, STREAM, WEEK, WINDOW, SHOWING, STATUS = "B4", "C4", "E4", "H4", "I4", "K4", "L4", "B5"
MAX_MATCHES = 250
WEEK_LIST_LEN = 104


def _sheet(wb, name):
    if name in wb.sheetnames:
        return wb[name]
    ws = wb.create_sheet(name)
    ws.sheet_state = "hidden"
    return ws


def search_helper(wb, key: str, dash: str, block: int, n_rows: int) -> str:
    """Build a 4-column block in CALC_SEARCH and return the defined name of the list."""
    ws = _sheet(wb, "CALC_SEARCH")
    c0 = block * 5 + 1
    A, B, C, D = (col(c0 + i) for i in range(4))
    scope, term = abs_ref(dash, SCOPE), abs_ref(dash, SEARCH)
    ws[f"{A}1"] = f"{key} match row"
    ws[f"{B}1"] = f"{key} k-th match"
    ws[f"{C}1"] = f"{key} name"
    ws[f"{D}1"] = f"{key} matches"
    last = n_rows + 1
    ws[f"{D}2"] = f"=COUNT({A}2:{A}{last})"
    ws[f"{D}3"] = "=COUNTA(MASTER_AREAS!$B:$B)-1"
    ws[f"{D}4"] = "=COUNTA(MASTER_ZONES!$B:$B)-1"
    for i in range(1, n_rows + 1):
        r = i + 1
        ws[f"{A}{r}"] = (
            f'=IF({scope}="Area",IF({i}<=${D}$3,IF(OR({term}="",ISNUMBER(SEARCH({term},INDEX(MASTER_AREAS!$B:$B,{i + 1})))),{i},""),""),'
            f'IF({scope}="Zone",IF({i}<=${D}$4,IF(OR({term}="",ISNUMBER(SEARCH({term},INDEX(MASTER_ZONES!$B:$B,{i + 1})))),{i},""),""),""))')
    for k in range(1, MAX_MATCHES + 1):
        r = k + 1
        ws[f"{B}{r}"] = f'=IFERROR(SMALL(${A}$2:${A}${last},{k}),"")'
        ws[f"{C}{r}"] = (f'=IF({B}{r}="","",IF({scope}="Area",INDEX(MASTER_AREAS!$B:$B,{B}{r}+1),'
                         f'INDEX(MASTER_ZONES!$B:$B,{B}{r}+1)))')
    name = f"{key}_SelectList"
    define(wb, name, f"OFFSET(CALC_SEARCH!${C}$2,0,0,MAX(1,MIN({MAX_MATCHES},CALC_SEARCH!${D}$2)),1)")
    return name


def week_list(wb, key: str, week_col_ref: str, block: int) -> tuple[str, str, str]:
    """'Latest' followed by the last 104 weeks that fall inside the data range.

    Returns (defined name of the list, latest-week cell, earliest-week cell)."""
    ws = _sheet(wb, "CALC_LISTS")
    c = col(block * 2 + 1)
    helper = col(block * 2 + 2)
    ws[f"{helper}1"] = f"{key} latest"
    ws[f"{helper}2"] = f"=MAX({week_col_ref})"
    ws[f"{helper}3"] = f"{key} earliest"
    ws[f"{helper}4"] = f"=MIN({week_col_ref})"
    ws[f"{helper}5"] = f"{key} weeks listed"
    ws[f"{helper}6"] = f"=COUNT({c}2:{c}{WEEK_LIST_LEN + 1})"
    ws[f"{c}1"] = "Latest"
    for k in range(WEEK_LIST_LEN):
        r = k + 2
        ws[f"{c}{r}"] = f'=IF(${helper}$2=0,"",IF(${helper}$2-{7 * k}>=${helper}$4,${helper}$2-{7 * k},""))'
        ws[f"{c}{r}"].number_format = "dd-mmm-yyyy"
    ws.column_dimensions[c].width = 13
    name = f"{key}_WeekList"
    define(wb, name, f"OFFSET(CALC_LISTS!${c}$1,0,0,1+CALC_LISTS!${helper}$6,1)")
    return name, f"CALC_LISTS!${helper}$2", f"CALC_LISTS!${helper}$4"


def control_strip(wb, ws, key: str, colour: str, stream_label: str, select_list: str, week_list_name: str,
                  last_col: str = "O", stream_options: str | None = None):
    label_f = font(8, True, tint(colour, 0.25))
    input_bg = "FFFFFF"
    strip_bg = tint(colour, 0.9)
    for r in (3, 4, 5):
        for c in range(2, ord(last_col) - 64 + 1):
            ws.cell(r, c).fill = fill(strip_bg)
    ws.row_dimensions[3].height = 16
    ws.row_dimensions[4].height = 24
    ws.row_dimensions[5].height = 18
    labels = [("B3", "SCOPE", None), ("C3", "SEARCH  (type part of a name)", "D3"),
              ("E3", "SELECT AREA / ZONE  ▼", "G3"), ("H3", "STREAM", None), ("I3", "WEEK ENDING  ▼", "J3"),
              ("K3", "COMPARE (WEEKS)", None), ("L3", "SHOWING", f"{last_col}3")]
    for ref, text, to in labels:
        put(ws, ref, text, f=label_f, align=LEFT, merge_to=to)
    inp = font(11, True)
    put(ws, SCOPE, "Global", f=inp, bg=input_bg, border=BOX, align=CENTER)
    put(ws, SEARCH, None, f=inp, bg=input_bg, border=BOX, align=LEFT, merge_to="D4")
    put(ws, SELECT, None, f=inp, bg=input_bg, border=BOX, align=LEFT, merge_to="G4")
    put(ws, STREAM, stream_label, f=inp, bg=input_bg if stream_options else tint(colour, 0.8), border=BOX,
        align=CENTER)
    put(ws, WEEK, "Latest", f=inp, bg=input_bg, border=BOX, align=CENTER, fmt="dd-mmm-yyyy", merge_to="J4")
    put(ws, WINDOW, 4, f=inp, bg=input_bg, border=BOX, align=CENTER)
    ws[SHOWING].alignment = LEFT

    dv_scope = DataValidation(type="list", formula1='"Global,Area,Zone"', allow_blank=False)
    dv_scope.error, dv_scope.errorTitle = "Choose Global, Area or Zone.", "Scope"
    dv_select = DataValidation(type="list", formula1=f"={select_list}", allow_blank=True)
    dv_select.error = "Pick a name from the list. Type part of the name in SEARCH to narrow it."
    dv_select.errorTitle = "Select"
    dv_week = DataValidation(type="list", formula1=f"={week_list_name}", allow_blank=True)
    dv_week.error, dv_week.errorTitle = "Pick 'Latest' or a week from the list.", "Week ending"
    dv_win = DataValidation(type="list", formula1='"1,2,4,8,13,26,52"', allow_blank=True)
    dv_win.error, dv_win.errorTitle = "Choose 1, 2, 4, 8, 13, 26 or 52.", "Comparison window"
    for dv, ref in ((dv_scope, SCOPE), (dv_select, SELECT), (dv_week, WEEK), (dv_win, WINDOW)):
        dv.showErrorMessage = True
        ws.add_data_validation(dv)
        dv.add(ref)
    if stream_options:
        dv_stream = DataValidation(type="list", formula1=f'"{stream_options}"', allow_blank=False)
        dv_stream.showErrorMessage = True
        ws.add_data_validation(dv_stream)
        dv_stream.add(STREAM)
    ws.freeze_panes = "A6"
