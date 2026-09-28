"""Data tables, master lists, Calendar and SETTINGS.

Every data sheet is one continuous Excel Table.  Computed columns
(Week_Ending, Adult_Total, Grand_Total, ...) are formulas, so the workbook
never trusts totals from incoming files.
"""
from __future__ import annotations

import datetime as dt
from dataclasses import dataclass, field

import pandas as pd
from openpyxl.worksheet.table import Table, TableStyleInfo

from config import COLOURS
from .common import (FMT_DATE, FMT_DATETIME, FMT_INT0, FMT_NGN, FMT_PCT, INK, LEFT, LEFT_WRAP, MUTED, WHITE,
                     col, define, fill, font, put, tint)

WEEK_ENDING_F = '=IF({Report_Date}{r}="","",{Report_Date}{r}+MOD(WeekEndDay-WEEKDAY({Report_Date}{r},2),7))'
ADULT_F = '=IF({Male}{r}="","",{Male}{r}+{Female}{r})'
GRAND_F = '=IF({Male}{r}="","",{Adult_Total}{r}+{Children}{r})'


@dataclass
class Column:
    name: str
    source: str | None = None       # dataframe column, or None for a formula column
    formula: str | None = None      # template using {Header} letters and {r}
    fmt: str | None = None
    width: float = 12
    kind: str = "text"              # text | date | datetime | int | num


@dataclass
class TableSpec:
    sheet: str
    table: str
    columns: list[Column]
    letters: dict = field(init=False)

    def __post_init__(self):
        self.letters = {c.name: col(i + 1) for i, c in enumerate(self.columns)}

    def ref(self, name: str) -> str:
        """Whole-column reference, e.g. DATA_WSF!$E:$E."""
        L = self.letters[name]
        return f"{self.sheet}!${L}:${L}"


def attendance_cols(unit_cols: list[Column]) -> list[Column]:
    return [Column("Week_Ending", formula=WEEK_ENDING_F, fmt=FMT_DATE, width=13, kind="date"),
            Column("Report_Date", "report_date", fmt=FMT_DATE, width=13, kind="date"),
            *unit_cols,
            Column("Male", "male", fmt=FMT_INT0, width=9, kind="int"),
            Column("Female", "female", fmt=FMT_INT0, width=9, kind="int"),
            Column("Adult_Total", formula=ADULT_F, fmt=FMT_INT0, width=11),
            Column("Children", "children", fmt=FMT_INT0, width=9, kind="int"),
            Column("Grand_Total", formula=GRAND_F, fmt=FMT_INT0, width=11)]


AUDIT = [Column("Source_File", "source_file", width=34), Column("Loaded_On", "loaded_on", fmt=FMT_DATETIME, width=17,
                                                                 kind="datetime")]

SPECS = {
    "WSF": TableSpec("DATA_WSF", "tblWSF", [
        Column("Week_Ending", formula=WEEK_ENDING_F, fmt=FMT_DATE, width=13, kind="date"),
        Column("Report_Date", "report_date", fmt=FMT_DATE, width=13, kind="date"),
        Column("Area", "area", width=18), Column("Zone", "zone", width=24),
        Column("Cells_Total", "cells_total", fmt=FMT_INT0, width=11, kind="int"),
        Column("Cells_Reported", "cells_reported", fmt=FMT_INT0, width=13, kind="int"),
        Column("Cells_Not_Reported", formula='=IF({Cells_Total}{r}="","",{Cells_Total}{r}-{Cells_Reported}{r})',
               fmt=FMT_INT0, width=16),
        Column("Male", "male", fmt=FMT_INT0, width=9, kind="int"),
        Column("Female", "female", fmt=FMT_INT0, width=9, kind="int"),
        Column("Adult_Total", formula=ADULT_F, fmt=FMT_INT0, width=11),
        Column("Children", "children", fmt=FMT_INT0, width=9, kind="int"),
        Column("Grand_Total", formula=GRAND_F, fmt=FMT_INT0, width=11),
        *AUDIT]),
    "MIDWEEK": TableSpec("DATA_MIDWEEK", "tblMidweek",
                         attendance_cols([Column("Area", "area", width=18), Column("Zone", "zone", width=24)]) + AUDIT),
    "CHOP": TableSpec("DATA_CHOP", "tblCHOP",
                      attendance_cols([Column("Area", "area", width=18), Column("Zone", "zone", width=24)]) + AUDIT),
    "COMMUNITY": TableSpec("DATA_COMMUNITY", "tblCommunity",
                           attendance_cols([Column("Church", "church", width=30),
                                            Column("Service_Type", "service_type", width=12)]) + AUDIT),
    "TRANSPORT_OPS": TableSpec("DATA_TRANSPORT_OPS", "tblTransportOps", [
        Column("Week_Ending", formula=WEEK_ENDING_F, fmt=FMT_DATE, width=13, kind="date"),
        Column("Report_Date", "report_date", fmt=FMT_DATE, width=13, kind="date"),
        Column("Category", "category", width=15), Column("Vehicle_ID", "vehicle_id", width=11),
        Column("Area", "area", width=18), Column("Zone", "zone", width=24),
        Column("Hired_By_Type", "hired_by_type", width=13),
        Column("Trips", "trips", fmt=FMT_INT0, width=8, kind="int"),
        Column("Ridership", "ridership", fmt=FMT_INT0, width=10, kind="int"),
        Column("Operational", "operational", width=11),
        Column("Capacity", formula='=IF({Vehicle_ID}{r}="","",IFERROR(INDEX(MASTER_FLEET!$C:$C,'
                                   'MATCH({Vehicle_ID}{r},MASTER_FLEET!$A:$A,0)),0))', fmt=FMT_INT0, width=9),
        Column("Seats_Offered", formula='=IF({Vehicle_ID}{r}="","",IF({Operational}{r}="Y",{Capacity}{r}*{Trips}{r},0))',
               fmt=FMT_INT0, width=12),
        *AUDIT]),
    "TRANSPORT_FINANCE": TableSpec("DATA_TRANSPORT_FINANCE", "tblTransportFinance", [
        Column("Week_Ending", formula=WEEK_ENDING_F, fmt=FMT_DATE, width=13, kind="date"),
        Column("Report_Date", "report_date", fmt=FMT_DATE, width=13, kind="date"),
        Column("Category", "category", width=15), Column("Area", "area", width=18), Column("Zone", "zone", width=24),
        Column("Cost_Type", "cost_type", width=16),
        Column("Amount", "amount", fmt=FMT_NGN, width=13, kind="num"),
        Column("Paid_By", "paid_by", width=11),
        *AUDIT]),
}

# Cell-level WSF detail for the latest weeks (feeds "which cells didn't report")
CELLS_CURRENT = TableSpec("WSF_CELLS_CURRENT", "tblWSFCellsCurrent", [
    Column("Week_Ending", formula=WEEK_ENDING_F, fmt=FMT_DATE, width=13, kind="date"),
    Column("Report_Date", "report_date", fmt=FMT_DATE, width=13, kind="date"),
    Column("Area", "area", width=18), Column("Zone", "zone", width=24), Column("Cell", "cell", width=15),
    Column("Reported", "reported", width=9),
    Column("Male", "male", fmt=FMT_INT0, width=8, kind="int"),
    Column("Female", "female", fmt=FMT_INT0, width=8, kind="int"),
    Column("Adult_Total", formula=ADULT_F, fmt=FMT_INT0, width=11),
    Column("Children", "children", fmt=FMT_INT0, width=9, kind="int"),
    Column("Grand_Total", formula=GRAND_F, fmt=FMT_INT0, width=11),
    *AUDIT,
    # helper: row number when this cell is a non-reporter inside the WSF dashboard's filter
    Column("List_Helper", formula='=IF(AND({Reported}{r}="N",{Week_Ending}{r}=CALC_WSF!$B$10,'
                                  'OR(CALC_WSF!$B$6="*",{Area}{r}=CALC_WSF!$B$6),'
                                  'OR(CALC_WSF!$B$7="*",{Zone}{r}=CALC_WSF!$B$7)),ROW(),"")', width=10),
])


def _to_date(v):
    if v is None or v == "" or (isinstance(v, float) and pd.isna(v)):
        return None
    if isinstance(v, (dt.date, dt.datetime)):
        return v
    return dt.date.fromisoformat(str(v)[:10])


def _to_datetime(v):
    if v is None or v == "" or (isinstance(v, float) and pd.isna(v)):
        return None
    return dt.datetime.fromisoformat(str(v))


def write_table(wb, spec: TableSpec, df: pd.DataFrame, colour: str, note: str | None = None):
    """Write a dataframe as a single continuous Excel Table (plus one blank row when empty)."""
    ws = wb.create_sheet(spec.sheet)
    ws.sheet_properties.tabColor = tint(colour, 0.35)
    headers = [c.name for c in spec.columns]
    ws.append(headers)
    converters = []
    for c in spec.columns:
        if c.formula:
            converters.append(None)
        elif c.kind == "date":
            converters.append(_to_date)
        elif c.kind == "datetime":
            converters.append(_to_datetime)
        elif c.kind == "int":
            converters.append(lambda v: None if v is None or pd.isna(v) else int(v))
        elif c.kind == "num":
            converters.append(lambda v: None if v is None or pd.isna(v) else float(v))
        else:
            converters.append(lambda v: "" if v is None or (isinstance(v, float) and pd.isna(v)) else v)
    src = [c.source for c in spec.columns]
    records = df[[s for s in src if s]].to_dict("records") if len(df) else []
    n = max(1, len(records))
    templates = [c.formula for c in spec.columns]
    L = spec.letters
    for i in range(n):
        r = i + 2
        rec = records[i] if records else {}
        row = []
        for s, conv, t in zip(src, converters, templates):
            if t:
                row.append(t.format(r=r, **L))
            else:
                row.append(conv(rec.get(s)) if rec else None)
        ws.append(row)
    for j, c in enumerate(spec.columns, start=1):
        ws.column_dimensions[col(j)].width = c.width
        if c.fmt:
            for (cell,) in ws.iter_rows(min_row=2, max_row=n + 1, min_col=j, max_col=j):
                cell.number_format = c.fmt
    for cell in ws[1]:
        cell.font = font(10, True, WHITE)
        cell.fill = fill(colour)
    last = col(len(spec.columns))
    t = Table(displayName=spec.table, ref=f"A1:{last}{n + 1}")
    t.tableStyleInfo = TableStyleInfo(name="TableStyleLight1", showRowStripes=True)
    ws.add_table(t)
    ws.freeze_panes = "A2"
    return ws


# --------------------------------------------------------------------------
# Masters
# --------------------------------------------------------------------------


def write_masters(wb, masters) -> dict[str, int]:
    colour = COLOURS["QUALITY"]
    cells = masters.cells.merge(masters.zones[["Zone_Name", "Area_Name"]], on="Zone_Name", how="left")
    frames = {
        "MASTER_AREAS": (masters.areas, "tblAreas", [14, 22, 8]),
        "MASTER_ZONES": (masters.zones, "tblZones", [12, 26, 22, 10, 8]),
        "MASTER_CELLS": (cells[["Cell", "Zone_Name", "Area_Name", "Operational"]], "tblCells", [16, 26, 22, 12]),
        "MASTER_COMMUNITY": (masters.community, "tblCommunityMaster", [10, 34, 18, 8]),
        "MASTER_FLEET": (masters.fleet, "tblFleet", [12, 16, 10, 10, 10, 20, 24, 14]),
    }
    sizes = {}
    for sheet, (df, tname, widths) in frames.items():
        ws = wb.create_sheet(sheet)
        ws.sheet_properties.tabColor = tint(colour, 0.5)
        ws.append(list(df.columns))
        rows = df.values.tolist()
        for r in rows or [[None] * len(df.columns)]:
            if sheet == "MASTER_FLEET" and r[2] not in ("", None):
                r = list(r)
                r[2] = int(float(r[2]))
            ws.append(r)
        for j, w in enumerate(widths, start=1):
            ws.column_dimensions[col(j)].width = w
        for cell in ws[1]:
            cell.font = font(10, True, WHITE)
            cell.fill = fill(colour)
        t = Table(displayName=tname, ref=f"A1:{col(len(df.columns))}{max(1, len(rows)) + 1}")
        t.tableStyleInfo = TableStyleInfo(name="TableStyleLight1", showRowStripes=True)
        ws.add_table(t)
        ws.freeze_panes = "A2"
        sizes[sheet] = len(rows)
    return sizes


# --------------------------------------------------------------------------
# SETTINGS and CALENDAR
# --------------------------------------------------------------------------

SETTINGS_ROWS = [
    # (name, value, format, note, editable)
    ("WeekEndDay", None, "0", "Day the reporting week ends: 1=Mon ... 6=Sat, 7=Sun. Change it in settings.yaml, "
                              "then run ingest.py so stored data and the workbook stay in step.", False),
    ("DefaultWindow", 4, "0", "Comparison window used when a dashboard's window box is empty (1, 2, 4, 8, 13, 26 or 52).", True),
    ("RAG_Green", 0.90, FMT_PCT, "Reporting / utilisation at or above this is green.", True),
    ("RAG_Amber", 0.75, FMT_PCT, "At or above this (but below green) is amber; below it is red.", True),
    ("UtilLow", 0.50, FMT_PCT, "Transport: buses/routes under this utilisation are flagged (X in the spec).", True),
    ("UtilHigh", 1.00, FMT_PCT, "Transport: utilisation above this means demand exceeds capacity.", True),
    ("UtilLowWeeks", 4, "0", "Transport: consecutive weeks under UtilLow before a plain-language flag is raised.", True),
    ("Cap_FT", 22, "0", "Seats per trip, FT Procured (LT) buses.", True),
    ("Cap_Coaster", 30, "0", "Seats per trip, church-owned Coasters (assumed 30; confirm).", True),
    ("Cap_Electric", 70, "0", "Riders per trip, electric buses.", True),
    ("Cap_Big", 70, "0", "Riders per trip, big buses.", True),
    ("CurrencySymbol", "₦", "@", "Currency shown on finance figures (NGN).", True),
]


def read_existing_settings(path) -> dict:
    """Keep values staff typed into SETTINGS when the workbook is rebuilt."""
    try:
        from openpyxl import load_workbook
        wb = load_workbook(path, read_only=True, data_only=True)
        if "SETTINGS" not in wb.sheetnames:
            return {}
        out = {}
        for row in wb["SETTINGS"].iter_rows(min_row=4, max_col=2, values_only=True):
            if row[0]:
                out[str(row[0])] = row[1]
        wb.close()
        return out
    except Exception:
        return {}


def write_settings(wb, settings: dict, existing: dict):
    ws = wb.create_sheet("SETTINGS")
    ws.sheet_properties.tabColor = "999999"
    put(ws, "A1", "SETTINGS", f=font(16, True, COLOURS["MASTER"]))
    put(ws, "A2", "Edit the blue values. They are kept when ingest.py rebuilds the workbook.",
        f=font(9, False, MUTED, italic=True))
    put(ws, "A3", "Setting", f=font(10, True, WHITE), bg=COLOURS["MASTER"])
    put(ws, "B3", "Value", f=font(10, True, WHITE), bg=COLOURS["MASTER"])
    put(ws, "C3", "What it does", f=font(10, True, WHITE), bg=COLOURS["MASTER"])
    for i, (name, default, fmt, note, editable) in enumerate(SETTINGS_ROWS, start=4):
        value = settings["week_end_day"] if name == "WeekEndDay" else existing.get(name, default)
        put(ws, f"A{i}", name, f=font(10, True))
        put(ws, f"B{i}", value, f=font(10, False, "0000FF" if editable else INK), fmt=fmt,
            bg="FFFFE0" if editable else None)
        put(ws, f"C{i}", note, f=font(9, False, MUTED), align=LEFT_WRAP)
        define(wb, name, f"SETTINGS!$B${i}")
    r = len(SETTINGS_ROWS) + 6
    put(ws, f"A{r}", "Stream colours (used on every sheet)", f=font(10, True))
    for j, (k, v) in enumerate(COLOURS.items(), start=r + 1):
        put(ws, f"A{j}", k, f=font(10))
        put(ws, f"B{j}", "#" + v, f=font(10, True, WHITE), bg=v)
    ws.column_dimensions["A"].width = 18
    ws.column_dimensions["B"].width = 12
    ws.column_dimensions["C"].width = 90
    ws.sheet_state = "hidden"


def write_calendar(wb, settings: dict):
    ws = wb.create_sheet("CALENDAR")
    ws.append(["Week_Ending", "Week_No", "Month", "Month_Name", "Quarter", "Year", "Year_Month"])
    start_year = settings["calendar_start_year"]
    n = 53 * settings["calendar_years"]
    ws.append([f"=DATE({start_year},1,1)+MOD(WeekEndDay-WEEKDAY(DATE({start_year},1,1),2),7)",
               "=INT((A2-DATE(YEAR(A2),1,1))/7)+1", "=MONTH(A2)", '=TEXT(A2,"mmm")', "=ROUNDUP(C2/3,0)",
               "=YEAR(A2)", '=F2&"-"&TEXT(C2,"00")'])
    for r in range(3, n + 2):
        ws.append([f"=A{r - 1}+7", f"=INT((A{r}-DATE(YEAR(A{r}),1,1))/7)+1", f"=MONTH(A{r})", f'=TEXT(A{r},"mmm")',
                   f"=ROUNDUP(C{r}/3,0)", f"=YEAR(A{r})", f'=F{r}&"-"&TEXT(C{r},"00")'])
    for r in range(2, n + 2):
        ws[f"A{r}"].number_format = FMT_DATE
    for c, w in zip("ABCDEFG", (13, 9, 7, 11, 8, 7, 11)):
        ws.column_dimensions[c].width = w
    t = Table(displayName="tblCalendar", ref=f"A1:G{n + 1}")
    t.tableStyleInfo = TableStyleInfo(name="TableStyleLight1", showRowStripes=True)
    ws.add_table(t)
    ws.sheet_state = "hidden"


# --------------------------------------------------------------------------
# DATA QUALITY
# --------------------------------------------------------------------------


def write_data_quality(wb, dq: pd.DataFrame, coverage: pd.DataFrame, missing: pd.DataFrame):
    colour = COLOURS["QUALITY"]
    from .common import dashboard_canvas, section_title, BOX, CENTER
    ws = wb.create_sheet("DATA QUALITY")
    ws.sheet_properties.tabColor = colour
    dashboard_canvas(ws, "QUALITY", "DATA QUALITY",
                     "Rejected rows, mismatched totals, duplicates, unknown Areas/Zones/Cells and missing weeks. "
                     "Fix the source file and drop it in the inbox again; a re-sent file replaces the old rows.",
                     n_cols=9, widths={"B": 20, "C": 11, "D": 34, "E": 26, "F": 7, "G": 12, "H": 60, "I": 17})
    ws.freeze_panes = "A4"

    section_title(ws, 4, "Issues by type", colour, last="I")
    put(ws, "B5", "Stream", f=font(9, True, WHITE), bg=colour)
    put(ws, "C5", "Severity", f=font(9, True, WHITE), bg=colour)
    put(ws, "D5", "Issue", f=font(9, True, WHITE), bg=colour)
    put(ws, "E5", "Rows", f=font(9, True, WHITE), bg=colour)
    log_first = 0
    summary = (dq.groupby(["stream", "severity", "issue"]).size().reset_index(name="n")
               if len(dq) else pd.DataFrame(columns=["stream", "severity", "issue", "n"]))
    r = 6
    # summary rows count the log below with COUNTIFS so it stays consistent with the table
    summary_rows = summary.values.tolist() or [["–", "–", "No issues logged", 0]]
    log_start_row = r + len(summary_rows) + 3 + max(len(coverage), 1) + 3 + max(len(missing), 1) + 3
    log_first = log_start_row + 1
    for stream, sev, issue, _ in summary_rows:
        put(ws, f"B{r}", stream, f=font(9))
        put(ws, f"C{r}", sev, f=font(9, True, {"Rejected": "A61B1B", "Warning": "8A5A00"}.get(sev, MUTED)))
        put(ws, f"D{r}", issue, f=font(9))
        if issue == "No issues logged":
            put(ws, f"E{r}", 0, f=font(9), fmt=FMT_INT0)
        else:
            put(ws, f"E{r}", f'=COUNTIFS($B${log_first}:$B$1048576,B{r},$C${log_first}:$C$1048576,C{r},'
                             f'$D${log_first}:$D$1048576,D{r})', f=font(9, True), fmt=FMT_INT0)
        r += 1

    r += 1
    section_title(ws, r, "Stream coverage", colour, last="I")
    r += 1
    for j, h in enumerate(["Stream", "First week", "Latest week", "Weeks loaded", "Rows in workbook"]):
        put(ws, f"{col(2 + j)}{r}", h, f=font(9, True, WHITE), bg=colour)
    r += 1
    for rec in coverage.to_dict("records") or [{"stream": "–"}]:
        put(ws, f"B{r}", rec.get("stream"), f=font(9))
        put(ws, f"C{r}", _to_date(rec.get("first")), f=font(9), fmt=FMT_DATE)
        put(ws, f"D{r}", _to_date(rec.get("latest")), f=font(9), fmt=FMT_DATE)
        put(ws, f"E{r}", rec.get("weeks"), f=font(9), fmt=FMT_INT0)
        put(ws, f"F{r}", rec.get("rows"), f=font(9), fmt=FMT_INT0)
        r += 1

    r += 1
    section_title(ws, r, "Missing weeks (no rows at all between a stream's first and latest week)", colour, last="I")
    r += 1
    put(ws, f"B{r}", "Stream", f=font(9, True, WHITE), bg=colour)
    put(ws, f"C{r}", "Week ending", f=font(9, True, WHITE), bg=colour)
    r += 1
    if len(missing):
        for rec in missing.to_dict("records"):
            put(ws, f"B{r}", rec["stream"], f=font(9))
            put(ws, f"C{r}", _to_date(rec["week_ending"]), f=font(9, True, "A61B1B"), fmt=FMT_DATE)
            r += 1
    else:
        put(ws, f"B{r}", "None — every stream has data for every week in its range.", f=font(9, False, MUTED, True))
        r += 1

    r = log_start_row - 1
    section_title(ws, r, "Issue log (newest first)", colour, last="I")
    headers = ["Stream", "Severity", "Issue", "Source file", "Row", "Week ending", "Detail", "Logged on"]
    for j, h in enumerate(headers):
        put(ws, f"{col(2 + j)}{log_start_row}", h, f=font(9, True, WHITE), bg=colour)
    rr = log_first
    for rec in dq.to_dict("records"):
        ws.cell(rr, 2, rec["stream"])
        ws.cell(rr, 3, rec["severity"])
        ws.cell(rr, 4, rec["issue"])
        ws.cell(rr, 5, rec["source_file"])
        ws.cell(rr, 6, rec["source_row"])
        ws.cell(rr, 7, _to_date(rec["week_ending"])).number_format = FMT_DATE
        ws.cell(rr, 8, rec["detail"])
        ws.cell(rr, 9, _to_datetime(rec["logged_on"])).number_format = FMT_DATETIME
        rr += 1
    if len(dq):
        t = Table(displayName="tblIssues", ref=f"B{log_start_row}:I{rr - 1}")
        t.tableStyleInfo = TableStyleInfo(name="TableStyleLight15", showRowStripes=True)
        ws.add_table(t)
    else:
        put(ws, f"B{log_first}", "No issues logged.", f=font(9, False, MUTED, True))
    ws.print_title_rows = "1:2"
