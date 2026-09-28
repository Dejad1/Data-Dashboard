"""Styles and small layout helpers shared by every sheet builder."""
from __future__ import annotations

from openpyxl.formatting.rule import FormulaRule
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter
from openpyxl.workbook.defined_name import DefinedName

from config import COLOURS

FONT = "Arial"

# Number formats.  Missing data is shown as the text "–" by the formulas, so a real zero stays 0.
FMT_INT = "#,##0"
FMT_INT0 = "#,##0"
FMT_DEC1 = "#,##0.0"
FMT_PCT = "0.0%"
FMT_PCT_SIGNED = '+0.0%;-0.0%;0.0%'
FMT_NGN = '"₦"#,##0'
FMT_NGN2 = '"₦"#,##0.00'
FMT_DATE = "dd-mmm-yyyy"
FMT_DATETIME = "yyyy-mm-dd hh:mm"
FMT_RATIO = '0.0" : 1"'

WHITE = "FFFFFF"
INK = "1F2933"
MUTED = "6B7785"
LINE = "D5DAE1"
PANEL = "F4F6F9"
GREEN_BG, GREEN_FG = "D8F0DF", "1E7B34"
AMBER_BG, AMBER_FG = "FFF1CC", "8A5A00"
RED_BG, RED_FG = "F9D9D9", "A61B1B"


def tint(hex_colour: str, amount: float) -> str:
    """Mix a colour with white; amount 0 = original, 1 = white."""
    r, g, b = (int(hex_colour[i:i + 2], 16) for i in (0, 2, 4))
    mix = lambda c: int(c + (255 - c) * amount)
    return f"{mix(r):02X}{mix(g):02X}{mix(b):02X}"


def font(size=10, bold=False, colour=INK, italic=False) -> Font:
    return Font(name=FONT, size=size, bold=bold, color=colour, italic=italic)


def fill(hex_colour: str) -> PatternFill:
    return PatternFill("solid", start_color=hex_colour, end_color=hex_colour)


def side(colour=LINE, style="thin") -> Side:
    return Side(style=style, color=colour)


BOX = Border(left=side(), right=side(), top=side(), bottom=side())
BOTTOM = Border(bottom=side())
CENTER = Alignment(horizontal="center", vertical="center", wrap_text=True)
LEFT = Alignment(horizontal="left", vertical="center", wrap_text=False)
LEFT_WRAP = Alignment(horizontal="left", vertical="top", wrap_text=True)
RIGHT = Alignment(horizontal="right", vertical="center")


def put(ws, ref: str, value=None, *, f=None, bg=None, fmt=None, align=None, border=None, merge_to=None):
    """Write one cell (optionally a merged block) with styling in a single call."""
    c = ws[ref]
    if value is not None:
        c.value = value
    if f is not None:
        c.font = f
    if bg is not None:
        c.fill = fill(bg)
    if fmt is not None:
        c.number_format = fmt
    if align is not None:
        c.alignment = align
    if merge_to:
        ws.merge_cells(f"{ref}:{merge_to}")
        if border is not None or bg is not None:
            for row in ws[f"{ref}:{merge_to}"]:
                for cell in row:
                    if border is not None:
                        cell.border = border
                    if bg is not None:
                        cell.fill = fill(bg)
    elif border is not None:
        c.border = border
    return c


def define(wb, name: str, ref: str) -> None:
    wb.defined_names[name] = DefinedName(name, attr_text=ref)


def abs_ref(sheet: str, cell: str) -> str:
    col = "".join(ch for ch in cell if ch.isalpha())
    row = "".join(ch for ch in cell if ch.isdigit())
    return f"{q(sheet)}!${col}${row}"


def q(sheet: str) -> str:
    """Quote a sheet name for use in a formula."""
    return f"'{sheet}'" if any(ch in sheet for ch in " -()&") else sheet


def col(n: int) -> str:
    return get_column_letter(n)


def dashboard_canvas(ws, stream_key: str, title: str, subtitle: str, n_cols: int = 14, widths=None):
    """Blank dashboard page: no gridlines, title band, column widths, print setup."""
    colour = COLOURS[stream_key]
    ws.sheet_view.showGridLines = False
    ws.sheet_view.zoomScale = 90
    ws.column_dimensions["A"].width = 2
    for i in range(2, n_cols + 2):
        ws.column_dimensions[col(i)].width = (widths or {}).get(col(i), 14)
    ws.column_dimensions[col(n_cols + 2)].width = 2
    last = col(n_cols)
    ws.row_dimensions[1].height = 34
    for c in range(1, n_cols + 3):
        ws.cell(1, c).fill = fill(colour)
        ws.cell(2, c).fill = fill(colour)
    put(ws, "B1", title, f=font(18, True, WHITE), align=LEFT)
    put(ws, "B2", subtitle, f=font(9, False, tint(colour, 0.7)), align=LEFT)
    home = ws.cell(1, n_cols)
    home.value = "⌂ HOME"
    home.hyperlink = "#HOME!A1"
    home.font = font(10, True, WHITE)
    home.alignment = RIGHT
    ws.page_setup.orientation = "landscape"
    ws.page_setup.paperSize = ws.PAPERSIZE_A4
    ws.page_setup.fitToWidth = 1
    ws.page_setup.fitToHeight = 0
    ws.sheet_properties.pageSetUpPr.fitToPage = True
    ws.print_options.horizontalCentered = True
    ws.page_margins.left = ws.page_margins.right = 0.3
    ws.page_margins.top = ws.page_margins.bottom = 0.4
    return colour, last


def section_title(ws, row: int, text: str, colour: str, first="B", last="N", note=None):
    put(ws, f"{first}{row}", text, f=font(12, True, colour), align=LEFT)
    for c in ws[f"{first}{row}:{last}{row}"][0]:
        c.border = Border(bottom=side(colour, "medium"))
    if note:
        put(ws, f"{last}{row}", note, f=font(8, False, MUTED, italic=True), align=RIGHT)
    ws.row_dimensions[row].height = 22


def rag_rules(ws, rng: str, first_cell: str):
    """Green / amber / red backgrounds on a % range, thresholds from SETTINGS."""
    fc = first_cell
    ws.conditional_formatting.add(rng, FormulaRule(
        formula=[f"AND(ISNUMBER({fc}),{fc}>=RAG_Green)"], fill=fill(GREEN_BG), font=Font(color=GREEN_FG, bold=True)))
    ws.conditional_formatting.add(rng, FormulaRule(
        formula=[f"AND(ISNUMBER({fc}),{fc}>=RAG_Amber,{fc}<RAG_Green)"], fill=fill(AMBER_BG),
        font=Font(color=AMBER_FG, bold=True)))
    ws.conditional_formatting.add(rng, FormulaRule(
        formula=[f"AND(ISNUMBER({fc}),{fc}<RAG_Amber)"], fill=fill(RED_BG), font=Font(color=RED_FG, bold=True)))


def arrow_rules(ws, rng: str, first_cell: str, good_up=True):
    up, down = (GREEN_FG, RED_FG) if good_up else (RED_FG, GREEN_FG)
    ws.conditional_formatting.add(rng, FormulaRule(formula=[f'LEFT({first_cell},1)="▲"'], font=Font(color=up, bold=True)))
    ws.conditional_formatting.add(rng, FormulaRule(formula=[f'LEFT({first_cell},1)="▼"'], font=Font(color=down, bold=True)))


def signed_rules(ws, rng: str, first_cell: str):
    ws.conditional_formatting.add(rng, FormulaRule(formula=[f"AND(ISNUMBER({first_cell}),{first_cell}>0)"],
                                                   font=Font(color=GREEN_FG)))
    ws.conditional_formatting.add(rng, FormulaRule(formula=[f"AND(ISNUMBER({first_cell}),{first_cell}<0)"],
                                                   font=Font(color=RED_FG)))
